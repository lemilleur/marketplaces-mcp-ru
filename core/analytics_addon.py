"""Additional marketplace analytics, registered by core.combined.

Uses active cabinets with environment fallback. Daily snapshots are opt-in:
python -m core.analytics_addon snapshot
Set SNAPSHOT_DB_PATH to persistent storage in containers.
XLSX reports require optional openpyxl; CSV needs no additional dependency.
API payloads originate from the supplied analytics_addon.py and need live validation.
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import re
import sqlite3
import sys
import time
from datetime import datetime, timedelta, timezone
from typing import Any, Awaitable, Callable

import httpx

# --------------------------------------------------------------------------- настройки

OZON_SELLER = "https://api-seller.ozon.ru"
OZON_PERF = "https://api-performance.ozon.ru"
WB_CONTENT = "https://content-api.wildberries.ru"
WB_ANALYTICS = "https://seller-analytics-api.wildberries.ru"
WB_PRICES = "https://discounts-prices-api.wildberries.ru"

DB_PATH = os.getenv("SNAPSHOT_DB_PATH", os.path.join(os.getenv("MARKETPLACE_MCP_HOME", os.path.expanduser("~/.marketplace-mcp")), "snapshots.sqlite"))
HTTP_TIMEOUT = float(os.getenv("ADDON_HTTP_TIMEOUT", "60"))
WB_FUNNEL_BATCH = max(1, min(20, int(os.getenv("WB_FUNNEL_BATCH", "20"))))          # nmId в одном запросе
WB_FUNNEL_PAUSE_SEC = float(os.getenv("WB_FUNNEL_PAUSE_SEC", "20"))  # пауза между пачками
MSK = timezone(timedelta(hours=3))

# Метрики воронки, на которые Ozon отвечает 400 «deprecated metrics used».
# Проверено запросами 21.10.2026 (первые пять) и по сообщениям разработчиков (остальные).
OZON_DEPRECATED_METRICS = {
    "hits_view_search", "hits_view_pdp", "hits_view",
    "session_view", "session_view_search", "session_view_pdp",
    "hits_tocart", "hits_tocart_search", "hits_tocart_pdp",
    "conv_tocart", "conv_tocart_search", "conv_tocart_pdp",
    "position_category",
}


def get_credentials() -> dict[str, str]:
    """Resolve credentials from the existing active cabinets, then environment."""
    from core.credentials import CredentialStore
    store = CredentialStore()
    result = {}
    for service, fields, prefix in (
        ("ozon", ["client_id", "api_key"], "OZON"),
        ("ozon_perf", ["client_id", "client_secret"], "OZON_PERF"),
        ("wb", ["token"], "WB_API"),
    ):
        creds, _ = store.resolve(service, fields, {f: f"{prefix}_{f.upper()}" for f in fields})
        result.update({f"{service}_{f}": v for f, v in creds.items()})
    return result


def _snapshot_scope(marketplace: str) -> str:
    # Separate histories even for env-only installs and cabinets with equal names.
    import hashlib
    c = get_credentials()
    keys = {"ozon": "ozon_client_id", "ozon_perf": "ozon_perf_client_id", "wb": "wb_token"}
    identity = c[keys[marketplace]]
    if not identity:
        raise ValueError(f"Missing credentials for {marketplace}")
    return marketplace + ":" + hashlib.sha256(identity.encode()).hexdigest()


# --------------------------------------------------------------------------- общее


def ok(data: Any, **extra: Any) -> dict:
    return {"ok": True, "data": data, **extra}


def fail(error: str, message: str, **extra: Any) -> dict:
    return {"ok": False, "error": error, "message": message, **extra}


class ApiError(Exception):
    def __init__(self, status: int, body: str, url: str):
        super().__init__(f"API request failed (HTTP {status})")
        self.status, self.body, self.url = status, body, url


async def _request(method: str, url: str, *, headers: dict | None = None,
                   json_body: Any = None, params: dict | None = None,
                   retries: int = 4) -> httpx.Response:
    """HTTP-запрос с повтором при 429 и 5xx (пауза растёт 2→4→8… до 60 с)."""
    delay = 2.0
    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT) as client:
        for attempt in range(retries + 1):
            r = await client.request(method, url, headers=headers, json=json_body, params=params)
            if (r.status_code == 429 or r.status_code >= 500) and attempt < retries:
                ra = r.headers.get("Retry-After") or r.headers.get("X-Ratelimit-Retry")
                try:
                    wait = float(ra) if ra else delay
                except ValueError:
                    wait = delay
                await asyncio.sleep(min(wait, 60))
                delay = min(delay * 2, 60)
                continue
            if r.status_code >= 400:
                raise ApiError(r.status_code, r.text, url)
            return r
    raise RuntimeError("unreachable")


def _ozon_headers() -> dict:
    c = get_credentials()
    return {"Client-Id": c["ozon_client_id"], "Api-Key": c["ozon_api_key"],
            "Content-Type": "application/json"}


def _wb_headers() -> dict:
    return {"Authorization": get_credentials()["wb_token"], "Content-Type": "application/json"}


async def _ozon_post(path: str, body: dict) -> dict:
    r = await _request("POST", OZON_SELLER + path, headers=_ozon_headers(), json_body=body)
    return r.json()


def _iso_day(d: str, end: bool = False) -> str:
    """'2026-10-14' -> '2026-10-14T00:00:00Z' (или 23:59:59 для конца периода)."""
    return f"{d}T23:59:59Z" if end else f"{d}T00:00:00Z"


def _to_number(s: str) -> Any:
    """'916,71' -> 916.71, '2166' -> 2166, прочее — как есть."""
    t = s.strip().replace("\u00a0", "").replace(" ", "")
    if re.fullmatch(r"-?\d+", t):
        return int(t)
    if re.fullmatch(r"-?\d+[.,]\d+", t):
        return float(t.replace(",", "."))
    return s.strip()


def _parse_delimited(text: str) -> list[dict]:
    """CSV с ';' или ',' -> список словарей; числа приводятся к числам."""
    text = text.lstrip("\ufeff")
    if not text.strip():
        return []
    first = text.splitlines()[0]
    delim = ";" if first.count(";") >= first.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    return [{k.strip(): _to_number(v or "") for k, v in row.items() if k} for row in reader]


def _parse_xlsx(content: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX reports require optional openpyxl dependency") from exc
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


# --------------------------------------------------------------------------- Ozon: аналитика


def map_analytics_rows(raw_rows: list[dict], dims: list[str], metrics: list[str]) -> tuple[list[dict], str | None]:
    """Превращает ответ /v1/analytics/data в плоские строки.
    Если Ozon вернул меньше метрик, чем запрошено, сопоставление по позициям ненадёжно —
    тогда отдаём сырые массивы и предупреждение."""
    out, warning = [], None
    for row in raw_rows:
        vals = row.get("metrics", [])
        item: dict[str, Any] = {}
        for i, d in enumerate(row.get("dimensions", [])):
            key = dims[i] if i < len(dims) else f"dim{i}"
            item[key] = d.get("id")
            if d.get("name"):
                item[f"{key}_name"] = d.get("name")
        if len(vals) == len(metrics):
            item.update(dict(zip(metrics, vals)))
        else:
            item["metrics_raw"] = vals
            warning = (f"Ozon вернул {len(vals)} метрик вместо {len(metrics)}: часть отброшена, "
                       "сопоставление по позициям ненадёжно — см. metrics_raw.")
        out.append(item)
    return out, warning


async def ozon_analytics_data_checked(date_from: str, date_to: str, metrics: list[str],
                                      dimension: list[str] | None = None) -> dict:
    """Аналитика Ozon /v1/analytics/data с пагинацией. Устаревшие метрики воронки
    убираются до запроса и перечисляются в dropped_metrics.
    Args: date_from/date_to 'YYYY-MM-DD'; metrics — напр. ['ordered_units','revenue'];
    dimension — напр. ['sku','day'] (по умолчанию ['day'])."""
    dims = dimension or ["day"]
    dropped = [m for m in metrics if m in OZON_DEPRECATED_METRICS]
    good = [m for m in metrics if m not in OZON_DEPRECATED_METRICS]
    if not good:
        return fail("deprecated_metrics", "Все запрошенные метрики отключены Ozon.",
                    dropped_metrics=dropped)
    rows, offset, limit = [], 0, 1000
    try:
        while True:
            resp = await _ozon_post("/v1/analytics/data", {
                "date_from": date_from, "date_to": date_to, "metrics": good,
                "dimension": dims, "limit": limit, "offset": offset,
                "sort": [{"key": dims[-1], "order": "ASC"}] if dims[-1] == "day" else [],
            })
            batch = resp.get("result", {}).get("data", [])
            rows.extend(batch)
            if len(batch) < limit:
                break
            offset += limit
    except ApiError as e:
        return fail("ozon_api", str(e), dropped_metrics=dropped)
    mapped, warning = map_analytics_rows(rows, dims, good)
    return ok(mapped, metrics=good, dropped_metrics=dropped, warning=warning)


async def ozon_product_queries_details(skus: list[str], date_from: str, date_to: str,
                                       limit_by_sku: int = 15, page_size: int = 100,
                                       sort_by: str = "BY_SEARCHES",
                                       sort_dir: str = "DESCENDING") -> dict:
    """Поисковые запросы, по которым находят товар (/v1/analytics/product-queries/details).
    Ограничения Ozon: последние 3 дня не отдаются (ещё считаются); данные старше месяца —
    только с Premium и только по неделям; часть полей без Premium пустая."""
    body = {"skus": [str(s) for s in skus], "date_from": _iso_day(date_from),
            "date_to": _iso_day(date_to, end=True), "limit_by_sku": limit_by_sku,
            "page": 0, "page_size": page_size, "sort_by": sort_by, "sort_dir": sort_dir}
    pages = []
    try:
        while True:
            resp = await _ozon_post("/v1/analytics/product-queries/details", body)
            pages.append(resp)
            page_count = resp.get("page_count")
            if not isinstance(page_count, int) or body["page"] + 1 >= page_count:
                break
            body["page"] += 1
    except ApiError as e:
        return fail("ozon_api", str(e))
    queries = []
    for p in pages:
        queries.extend(p.get("queries", []))
    return ok(queries if queries else pages,
              note="Пустые поля могут означать отсутствие Premium, а не ноль.")


# --------------------------------------------------------------------------- Ozon: отчёты

REPORT_PATH_RE = re.compile(r"^/v\d+/report/[a-z0-9_\-/]+/create$")


def report_preset(preset: str, date_from: str | None, date_to: str | None) -> tuple[str, dict]:
    """Готовые тела запросов. ПРОВЕРИТЬ поля фильтров по актуальной документации Ozon."""
    if preset == "products":
        return "/v1/report/products/create", {
            "language": "DEFAULT", "offer_id": [], "search": "", "sku": [], "visibility": "ALL"}
    if preset == "postings":
        if not (date_from and date_to):
            raise ValueError("Для postings нужны date_from и date_to")
        return "/v1/report/postings/create", {
            "filter": {"processed_at_from": _iso_day(date_from),
                       "processed_at_to": _iso_day(date_to, end=True),
                       "delivery_schema": ["fbo", "fbs", "rfbs"]},
            "language": "DEFAULT"}
    raise ValueError(f"Неизвестный preset: {preset}")


async def ozon_report(preset: str = "", path: str = "", body: dict | None = None,
                      date_from: str = "", date_to: str = "",
                      max_wait_sec: int = 300, max_rows: int = 5000) -> dict:
    """Отчёт Ozon целиком: создать -> ждать /v1/report/info -> скачать -> вернуть строки.
    Либо preset ('products' | 'postings'), либо path ('/v1/report/.../create') + body.
    Создание отчёта ничего не меняет в кабинете."""
    try:
        if preset:
            path, body = report_preset(preset, date_from or None, date_to or None)
        if not REPORT_PATH_RE.match(path or ""):
            return fail("invalid_params", "path должен иметь вид /vN/report/<тип>/create")
        created = await _ozon_post(path, body or {})
        code = (created.get("result") or {}).get("code") or created.get("code")
        if not code:
            return fail("ozon_api", f"Ozon не вернул код отчёта: {created}")
        deadline, info = time.monotonic() + max_wait_sec, {}
        while time.monotonic() < deadline:
            info = (await _ozon_post("/v1/report/info", {"code": code})).get("result", {})
            status = str(info.get("status", "")).lower()
            if status == "success" and info.get("file"):
                break
            if status in ("failed", "error"):
                return fail("report_failed", info.get("error") or "Ozon не смог собрать отчёт",
                            code=code)
            await asyncio.sleep(5)
        else:
            return fail("timeout", f"Отчёт не готов за {max_wait_sec} с", code=code)
        r = await _request("GET", info["file"])
        url = info["file"].lower().split("?")[0]
        if url.endswith(".xlsx") or r.content[:2] == b"PK":
            rows = _parse_xlsx(r.content)
        else:
            rows = _parse_delimited(r.content.decode("utf-8-sig", errors="replace"))
    except (ApiError, ValueError) as e:
        return fail("ozon_api", str(e))
    return ok(rows[:max_rows], total_rows=len(rows), truncated=len(rows) > max_rows,
              code=code, file=info.get("file"))


# --------------------------------------------------------------------------- Ozon Performance

_perf_token: dict[str, Any] = {"value": None, "expires": 0.0, "identity": None}


async def _perf_headers() -> dict:
    c = get_credentials()
    identity = (c["ozon_perf_client_id"], c["ozon_perf_client_secret"])
    if (_perf_token["identity"] != identity or not _perf_token["value"]
            or time.time() > _perf_token["expires"] - 60):
        r = await _request("POST", OZON_PERF + "/api/client/token", json_body={
            "client_id": c["ozon_perf_client_id"], "client_secret": c["ozon_perf_client_secret"],
            "grant_type": "client_credentials"})
        data = r.json()
        _perf_token["identity"] = identity
        _perf_token["value"] = data["access_token"]
        _perf_token["expires"] = time.time() + float(data.get("expires_in", 1800))
    return {"Authorization": f"Bearer {_perf_token['value']}"}


def summarize_perf_daily(rows: list[dict]) -> list[dict]:
    """Итоги по дням по всем кампаниям."""
    days: dict[str, dict] = {}
    for r in rows:
        d = days.setdefault(r.get("Дата"), {"date": r.get("Дата"), "views": 0, "clicks": 0,
                                            "spend": 0.0, "orders": 0, "orders_rub": 0.0})
        d["views"] += r.get("Показы") or 0
        d["clicks"] += r.get("Клики") or 0
        d["spend"] += r.get("Расход, ₽") or 0
        d["orders"] += r.get("Заказы, шт.") or 0
        d["orders_rub"] += r.get("Заказы, ₽") or 0
    out = []
    for d in sorted(days.values(), key=lambda x: str(x["date"])):
        d["spend"] = round(d["spend"], 2)
        d["drr_pct"] = round(d["spend"] / d["orders_rub"] * 100, 1) if d["orders_rub"] else None
        out.append(d)
    return out


async def ozon_perf_daily(date_from: str, date_to: str) -> dict:
    """Дневная статистика рекламы Ozon по кампаниям в JSON + итоги по дням с ДРР.
    Важно: Ozon относит рекламный заказ к дате клика, а не к дате заказа."""
    try:
        r = await _request("GET", OZON_PERF + "/api/client/statistics/daily",
                           headers=await _perf_headers(),
                           params={"dateFrom": date_from, "dateTo": date_to})
    except ApiError as e:
        return fail("ozon_perf_api", str(e))
    rows = _parse_delimited(r.text)
    return ok(rows, by_day=summarize_perf_daily(rows),
              note="Рекламный заказ привязан к дате клика.")


# --------------------------------------------------------------------------- Wildberries


async def wb_all_nm_ids() -> list[int]:
    """Все nmId продавца через список карточек (пагинация курсором)."""
    nm_ids: list[int] = []
    cursor: dict[str, Any] = {"limit": 100}
    while True:
        r = await _request("POST", WB_CONTENT + "/content/v2/get/cards/list", headers=_wb_headers(),
                           json_body={"settings": {"cursor": cursor, "filter": {"withPhoto": -1}}})
        data = r.json()
        cards = data.get("cards", [])
        nm_ids.extend(c["nmID"] for c in cards if "nmID" in c)
        cur = data.get("cursor", {})
        if len(cards) < cursor["limit"] or not cur.get("nmID"):
            break
        cursor = {"limit": 100, "updatedAt": cur.get("updatedAt"), "nmID": cur.get("nmID")}
    return nm_ids


async def wb_funnel_all(date_from: str, date_to: str, nm_ids: list[int] | None = None) -> dict:
    """Воронка WB по дням (переходы в карточку, корзина, заказы, выкупы, конверсии)
    по всем товарам. Сам собирает nmId, режет на пачки, соблюдает лимиты WB."""
    try:
        ids = nm_ids or await wb_all_nm_ids()
        if not ids:
            return fail("no_products", "У продавца не найдено карточек WB")
        result: list[Any] = []
        for i in range(0, len(ids), WB_FUNNEL_BATCH):
            if i:
                await asyncio.sleep(WB_FUNNEL_PAUSE_SEC)
            r = await _request("POST", WB_ANALYTICS + "/api/analytics/v3/sales-funnel/products/history",
                               headers=_wb_headers(), json_body={
                                   "selectedPeriod": {"start": date_from, "end": date_to},
                                   "nmIds": ids[i:i + WB_FUNNEL_BATCH],
                                   "skipDeletedNm": True, "aggregationLevel": "day"})
            data = r.json()
            batch = data.get("data", data) if isinstance(data, dict) else data
            result.extend(batch if isinstance(batch, list) else [batch])
    except ApiError as e:
        return fail("wb_api", str(e))
    return ok(result, nm_count=len(ids))


# --------------------------------------------------------------------------- снимки


def _db() -> sqlite3.Connection:
    os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.execute("""CREATE TABLE IF NOT EXISTS snapshots (
        snap_date TEXT NOT NULL, marketplace TEXT NOT NULL, kind TEXT NOT NULL,
        item_key TEXT NOT NULL, payload TEXT NOT NULL,
        PRIMARY KEY (snap_date, marketplace, kind, item_key))""")
    return con


def save_snapshot(snap_date: str, marketplace: str, kind: str,
                  items: list[dict], key_fn: Callable[[dict], Any]) -> int:
    marketplace = _snapshot_scope(marketplace)
    con = _db()
    with con:
        con.executemany(
            "INSERT OR REPLACE INTO snapshots VALUES (?,?,?,?,?)",
            [(snap_date, marketplace, kind, str(key_fn(it)), json.dumps(it, ensure_ascii=False))
             for it in items])
    con.close()
    return len(items)


async def _fetch_ozon_paged(path: str) -> list[dict]:
    items, cursor = [], ""
    while True:
        resp = await _ozon_post(path, {"cursor": cursor, "filter": {"visibility": "ALL"}, "limit": 1000})
        batch = resp.get("items", [])
        items.extend(batch)
        cursor = resp.get("cursor") or ""
        if not cursor or len(batch) < 1000:
            return items


async def _fetch_perf_campaigns() -> list[dict]:
    r = await _request("GET", OZON_PERF + "/api/client/campaign", headers=await _perf_headers())
    data = r.json()
    return data if isinstance(data, list) else data.get("list", [])


async def _fetch_wb_prices() -> list[dict]:
    items, offset = [], 0
    while True:
        r = await _request("GET", WB_PRICES + "/api/v2/list/goods/filter", headers=_wb_headers(),
                           params={"limit": 1000, "offset": offset})
        batch = (r.json().get("data") or {}).get("listGoods") or []
        items.extend(batch)
        if len(batch) < 1000:
            return items
        offset += 1000


SNAPSHOT_SOURCES: list[tuple[str, str, Callable[[], Awaitable[list[dict]]], Callable[[dict], Any]]] = [
    ("ozon", "prices", lambda: _fetch_ozon_paged("/v5/product/info/prices"), lambda it: it.get("offer_id")),
    ("ozon", "stocks", lambda: _fetch_ozon_paged("/v4/product/info/stocks"), lambda it: it.get("offer_id")),
    ("ozon_perf", "campaigns", _fetch_perf_campaigns, lambda it: it.get("id")),
    ("wb", "prices", _fetch_wb_prices, lambda it: it.get("nmID")),
]


async def snapshot_run(snap_date: str = "") -> dict:
    """Снять дневной снимок цен/акций, остатков и рекламных кампаний (Ozon, WB).
    Один упавший источник не останавливает остальные."""
    day = snap_date or datetime.now(MSK).strftime("%Y-%m-%d")
    report: dict[str, Any] = {"date": day, "saved": {}, "errors": {}}
    for mp, kind, fetch, key_fn in SNAPSHOT_SOURCES:
        try:
            report["saved"][f"{mp}.{kind}"] = save_snapshot(day, mp, kind, await fetch(), key_fn)
        except Exception as e:  # noqa: BLE001 — фиксируем любую ошибку источника
            report["errors"][f"{mp}.{kind}"] = str(e)[:300]
    return ok(report) if not report["errors"] else ok(report, warning="Часть источников не снята")


def _dig(obj: Any, dotted: str) -> Any:
    for part in dotted.split("."):
        obj = obj.get(part) if isinstance(obj, dict) else None
    return obj


async def history_get(marketplace: str, kind: str, date_from: str, date_to: str,
                      item_key: str = "", fields: list[str] | None = None) -> dict:
    """История из снимков по дням.
    Args: marketplace 'ozon'|'ozon_perf'|'wb'; kind 'prices'|'stocks'|'campaigns';
    item_key — offer_id (Ozon), nmID (WB), id кампании; пусто = все;
    fields — пути внутри записи, напр. ['price.marketing_seller_price','price.min_price']."""
    marketplace = _snapshot_scope(marketplace)
    con = _db()
    q = ("SELECT snap_date, item_key, payload FROM snapshots WHERE marketplace=? AND kind=? "
         "AND snap_date BETWEEN ? AND ?")
    args: list[Any] = [marketplace, kind, date_from, date_to]
    if item_key:
        q += " AND item_key=?"
        args.append(item_key)
    rows = con.execute(q + " ORDER BY snap_date, item_key", args).fetchall()
    con.close()
    out = []
    for d, key, payload in rows:
        rec = json.loads(payload)
        out.append({"date": d, "key": key,
                    **({f: _dig(rec, f) for f in fields} if fields else {"record": rec})})
    return ok(out, count=len(out))


# --------------------------------------------------------------------------- регистрация

TOOLS: list[tuple[str, Callable[..., Awaitable[dict]]]] = [
    ("ozon_analytics_data_checked", ozon_analytics_data_checked),
    ("ozon_product_queries_details", ozon_product_queries_details),
    ("ozon_report", ozon_report),
    ("ozon_perf_daily", ozon_perf_daily),
    ("wb_funnel_all", wb_funnel_all),
    ("snapshot_run", snapshot_run),
    ("history_get", history_get),
]


def register(mcp: Any) -> None:
    """Регистрирует инструменты на сервере (FastMCP v1, MCPServer v2, fastmcp)."""
    for name, fn in TOOLS:
        mcp.tool(name=name, description=(fn.__doc__ or "").strip())(fn)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "snapshot":
        print(json.dumps(asyncio.run(snapshot_run()), ensure_ascii=False, indent=2))
    else:
        print("Использование: python analytics_addon.py snapshot")
